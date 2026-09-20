#!/usr/bin/env python3
"""Assemble the FINAL research proposal (markdown files) into a single .docx (APA-style manuscript)."""
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

PARTS = [
    ("TITLE", None),
    ("Front-Matter-FINAL.md", "PRELIMINARY PAGES"),
    ("Chapter-1-Introduction-DRAFT.md", None),
    ("Chapter-2-Literature-Review-FINAL.md", None),
    ("Chapter-3-Research-Methodology-PILOT.md", None),
    ("Chapter-4-PILOT-Findings.md", None),
    ("Chapter-5-Summary-Conclusions-Framework.md", None),
    ("References-FINAL-APA7.md", "REFERENCES"),
    ("Pilot-Instruments-Print-Ready.md", "APPENDIX A: PILOT RESEARCH INSTRUMENTS"),
    ("Theory-of-Change.md", "APPENDIX B: THEORY OF CHANGE"),
    ("Gathered-Editions-Log.md", "APPENDIX C: EDITIONS LOG & DATASET NOTES"),
]

TITLE_BLOCK = """ASSESSING THE IMPACTS OF THE NEW VISION NEWSPAPER OVER THE PAST 40 YEARS OF ITS EXISTENCE (1986-2026): A MIXED-METHODS PILOT STUDY

RESEARCH PROPOSAL

BY

[FULL NAME]  -  [REG. NO.]

A Research Proposal Submitted to the [Department, Faculty] in Partial Fulfilment of the Requirements for the Award of [Degree] of [University]

[Month], 2026

Referencing style: APA 7th edition (in-text author-date; full reference list at the end). Content-analysis data are drawn from publicly archived editions of The New Vision (Wayback Machine 2007-2019; live site 2026); every coded item is traceable via Appendix C and the dataset file."""

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.5
for i in (1, 2, 3):
    hs = doc.styles[f"Heading {i}"]
    hs.font.name = "Times New Roman"
    hs.font.color.rgb = RGBColor(0, 0, 0)

def add_runs(para, text):
    text = text.replace("[ ]", "\u2610").replace("[x]", "\u2611").replace("[X]", "\u2611")
    for t in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`.+?`)", text):
        if not t:
            continue
        if t.startswith("**") and t.endswith("**") and len(t) > 4:
            r = para.add_run(t[2:-2]); r.bold = True
        elif t.startswith("*") and t.endswith("*") and len(t) > 2:
            r = para.add_run(t[1:-1]); r.italic = True
        elif t.startswith("`") and t.endswith("`") and len(t) > 2:
            r = para.add_run(t[1:-1]); r.font.name = "Courier New"; r.font.size = Pt(10)
        else:
            t = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", t)
            para.add_run(t)

def add_table(lines):
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
    rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = "Table Grid"
    for i, r in enumerate(rows):
        for j in range(ncols):
            p = tbl.cell(i, j).paragraphs[0]
            add_runs(p, r[j] if j < len(r) else "")
            for run in p.runs:
                run.font.size = Pt(10)
            if i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()

def process_md(path, override_title=None):
    with open(path) as f:
        lines = f.read().splitlines()
    if override_title:
        doc.add_heading(override_title, level=1)
    i, buf, skipped_h1 = 0, [], (not override_title)
    def flush_table():
        if buf:
            add_table(buf); buf.clear()
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("|") and "|" in s[1:]:
            buf.append(ln); i += 1; continue
        flush_table()
        if not s:
            i += 1; continue
        if s.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                r = doc.add_paragraph().add_run(lines[i]); r.font.name = "Courier New"; r.font.size = Pt(10)
                i += 1
            i += 1; continue
        if s in ("---", "***"):
            i += 1; continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            lvl = len(m.group(1))
            if lvl == 1 and not skipped_h1:
                skipped_h1 = True  # skip the file's own H1 when an override title is supplied
            else:
                doc.add_heading(re.sub(r"\*\*(.*?)\*\*", r"\1", m.group(2)), level=min(lvl, 3))
            i += 1; continue
        if s.startswith(">"):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(24)
            add_runs(p, s.lstrip("> ").strip())
            for r in p.runs: r.italic = True
            i += 1; continue
        m = re.match(r"^(\s*)[-*]\s+(.*)", ln)
        if m:
            add_runs(doc.add_paragraph(style="List Bullet"), m.group(2)); i += 1; continue
        m = re.match(r"^(\s*)\d+[.)]\s+(.*)", ln)
        if m:
            add_runs(doc.add_paragraph(style="List Number"), m.group(2)); i += 1; continue
        add_runs(doc.add_paragraph(), s); i += 1
    flush_table()

for fname, title in PARTS:
    if fname == "TITLE":
        for j, block in enumerate(TITLE_BLOCK.split("\n\n")):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(block); r.bold = (j == 0)
            r.font.size = Pt(16) if j == 0 else Pt(13)
        doc.add_page_break(); continue
    process_md(fname, title)
    doc.add_page_break()

out = "New-Vision-40-Years-RESEARCH-PROPOSAL-FINAL.docx"
doc.save(out)
print("saved", out)
