#!/usr/bin/env python3
"""Convert the project's markdown files to formatted Word (.docx) documents."""

import re
import os
import csv as csvmod

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INLINE_RE = re.compile(
    r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))'
)
LINK_RE = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')
BOLD_HEAD = re.compile(r'\*\*(.+?)\*\*')


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def add_inline(paragraph, text):
    """Add text to a paragraph, honouring **bold**, *italic*, `code`, [links](url)."""
    text = text.replace('<br>', '\n')
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            r.italic = True
        elif part.startswith('[') and '](' in part:
            m = LINK_RE.match(part)
            if m:
                r = paragraph.add_run(m.group(1))
                r.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
                r.underline = True
            else:
                paragraph.add_run(part)
        else:
            paragraph.add_run(part)


def split_row(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    return [c.strip() for c in line.split('|')]


def is_sep(line):
    s = line.strip()
    if not s.startswith('|'):
        return False
    return bool(re.fullmatch(r'\|[\s:\-|]+\|', s)) and '-' in s


def flush_table(doc, rows):
    """rows: list of lists of cell-strings. First row is the header."""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j in range(ncols):
            txt = row[j] if j < len(row) else ''
            cell = cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, txt)
            for r in p.runs:
                r.font.size = Pt(9.5)
                if i == 0:
                    r.bold = True
            if i == 0:
                shade(cell, 'E7EEF7')
    doc.add_paragraph()


def convert(md_path, docx_path, doc_title=None):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)

    table_buf = []
    in_code = False

    def flush():
        nonlocal table_buf
        flush_table(doc, table_buf)
        table_buf = []

    for raw in lines:
        line = raw.rstrip()

        # fenced code blocks
        if line.strip().startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            continue

        # tables
        if line.strip().startswith('|'):
            if is_sep(line):
                continue
            table_buf.append(split_row(line))
            continue
        else:
            if table_buf:
                flush()

        if not line.strip():
            continue

        # horizontal rule
        if re.fullmatch(r'-{3,}', line.strip()):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            bdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:color'), 'BFBFBF')
            bdr.append(bottom)
            pPr.append(bdr)
            continue

        # headings
        m = re.match(r'^(#{1,6})\s+(.*)$', line)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1:
                p = doc.add_heading('', level=0)
            else:
                p = doc.add_heading('', level=min(level - 1, 4))
            add_inline(p, text)
            continue

        # blockquote
        if line.strip().startswith('>'):
            txt = line.strip().lstrip('>').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, txt)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue

        # bullet list
        m = re.match(r'^(\s*)[-*+]\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, m.group(2))
            continue

        # numbered list
        m = re.match(r'^(\s*)\d+\.\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, m.group(2))
            continue

        # normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, line)

    if table_buf:
        flush()

    doc.save(docx_path)
    return docx_path


def csv_to_docx(csv_path, docx_path):
    with open(csv_path, encoding='utf-8') as f:
        rows = list(csvmod.reader(f))
    doc = Document()
    for s in doc.sections:
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)
    doc.add_heading('Men\'s Clothing Budget — Nairobi (spreadsheet view)', level=0)
    rows = [r for r in rows if any(c.strip() for c in r)]
    flush_table(doc, rows)
    doc.save(docx_path)
    return docx_path


if __name__ == '__main__':
    base = '/home/user/Gide26'
    jobs = [
        ('kenya-2-day-wholesale-clothes-trip.md', 'Kenya-2-Day-Wholesale-Clothes-Trip.docx'),
        ('men-clothes-budget-nairobi.md', 'Mens-Clothing-Budget-Nairobi.docx'),
        ('sourcing-directory-your-items.md', 'Sourcing-Directory-Your-Items.docx'),
    ]
    for md, dx in jobs:
        out = convert(os.path.join(base, md), os.path.join(base, dx))
        print('OK', os.path.basename(out), os.path.getsize(out), 'bytes')
    out = csv_to_docx(os.path.join(base, 'men-clothes-budget-nairobi.csv'),
                      os.path.join(base, 'Mens-Clothing-Budget-Spreadsheet.docx'))
    print('OK', os.path.basename(out), os.path.getsize(out), 'bytes')
